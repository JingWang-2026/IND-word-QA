from openpyxl import load_workbook
from sqlmodel import Session
from uuid import uuid4

from app.core.database import engine, init_db
from app.models.document import Document
from app.models.issue import QAIssue
from app.models.project import Project
from app.services.export_service import ExportService


def test_export_service_creates_xlsx(tmp_path) -> None:
    init_db()
    project = Project(id=f"project_{uuid4().hex}", name="Export Test")
    document = Document(
        id=f"doc_{uuid4().hex}",
        project_id=project.id,
        filename="Protocol.docx",
        stored_path="Protocol.docx",
        file_size=1,
        status="qa_done",
    )
    issue = QAIssue(
        id=f"QA-{uuid4().hex[:8].upper()}",
        project_id=project.id,
        document_id=document.id,
        severity="High",
        category="Numeric",
        rule_id="NUM-001",
        title="Percentage calculation mismatch",
        description="Mismatch.",
        source_text="3/20 (20.0%)",
        suggestion="Use 15.0%.",
    )
    with Session(engine) as session:
        session.add(project)
        session.add(document)
        session.add(issue)
        session.commit()
        content = ExportService().issues_xlsx(session, [issue])

    path = tmp_path / "issues.xlsx"
    path.write_bytes(content)
    workbook = load_workbook(path)
    sheet = workbook["QA Issues"]
    assert sheet["A1"].value == "Issue ID"
    assert sheet["E2"].value == "Protocol.docx"
    assert sheet["I2"].value == "3/20 (20.0%)"
